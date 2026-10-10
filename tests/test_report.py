"""Report tests, using the September 2026 reference figures as the golden fixture."""

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from fpdf import FPDF
from openpyxl import load_workbook
import pytest

from custom_components.plug_and_claim.report import build_report, format_kwh, write_pdf, write_xlsx

SEPTEMBER = {
    "2026-09-01": 6.09, "2026-09-02": 14.96, "2026-09-03": 8.425, "2026-09-04": 0.566,
    "2026-09-05": 7.014, "2026-09-06": 6.649, "2026-09-07": 2.933, "2026-09-08": 14.021,
    "2026-09-09": 9.558, "2026-09-10": 7.095504, "2026-09-11": 2.580446, "2026-09-12": 14.53855,
    "2026-09-13": 0, "2026-09-14": 5.899567, "2026-09-15": 3.449925, "2026-09-16": 16.391665,
    "2026-09-17": 7.139873, "2026-09-18": 9.721264, "2026-09-19": 17.068742, "2026-09-20": 10.059154,
    "2026-09-21": 5.873056, "2026-09-22": 10.698574, "2026-09-23": 9.31709, "2026-09-24": 12.98958,
    "2026-09-25": 0, "2026-09-26": 18.007821, "2026-09-27": 8.646209, "2026-09-28": 3.848412,
    "2026-09-29": 5.005205, "2026-09-30": 16.75645,
}
DETAILS = {
    "vehicle": "Mitsubishi Outlander PHEV",
    "registration": "ABC123",
    "source": "sensor.outlander_charger_today_s_consumption",
    "generated": datetime(2026, 10, 1, 1, 0),
}


def test_golden_totals() -> None:
    report = build_report("2026-09", SEPTEMBER, 0.407, **DETAILS)
    assert report.total_kwh == Decimal("255.303087")
    assert report.amount == Decimal("103.91")
    assert report.missing_days == []
    assert len(report.days) == 30
    assert report.label == "Sep-26"
    assert report.file_stem == "Energy_Usage_Sep_2026"


@pytest.mark.parametrize(
    ("kwh", "rate", "amount"),
    [
        ("255.303087", 0.407, "103.91"),
        ("10", 0.4125, "4.13"),  # exactly half a cent rounds up
        ("10", 0.4135, "4.14"),
        ("0", 0.407, "0.00"),
    ],
)
def test_amount_rounds_half_up(kwh: str, rate: float, amount: str) -> None:
    report = build_report("2026-02", {"2026-02-01": float(kwh)}, rate)
    assert report.amount == Decimal(amount)


def test_missing_days_are_listed_not_zero() -> None:
    days = {k: v for k, v in SEPTEMBER.items() if k not in ("2026-09-13", "2026-09-14")}
    report = build_report("2026-09", days, 0.407)
    assert report.missing_days == [date(2026, 9, 13), date(2026, 9, 14)]
    assert report.total_kwh == Decimal("249.403520")
    assert len(report.days) == 30


def test_month_lengths() -> None:
    assert len(build_report("2028-02", {}, 0.4).days) == 29  # leap year
    assert len(build_report("2026-10", {}, 0.4).days) == 31


def test_format_kwh() -> None:
    assert [format_kwh(Decimal(v)) for v in ("6.09", "0", "16.756450", "14.53855")] == [
        "6.09", "0", "16.75645", "14.53855",
    ]


def test_xlsx_golden(tmp_path: Path) -> None:
    report = build_report("2026-09", SEPTEMBER, 0.407, **DETAILS)
    path = tmp_path / f"{report.file_stem}.xlsx"
    write_xlsx(report, path)

    wb = load_workbook(path)
    assert wb.sheetnames == ["Sep-26"]
    ws = wb["Sep-26"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]

    summary = next(i for i, r in enumerate(rows) if r[1] == "Energy Usage(kWh)")
    assert rows[summary][2:4] == ["rate", "Amount"]
    period, total, rate, amount = rows[summary + 1][:4]
    assert period == datetime(2026, 9, 1)
    assert ws.cell(summary + 2, 1).number_format == "yyyy/mm"
    assert (total, rate, amount) == (255.303087, 0.407, 103.91)
    assert ws.cell(summary + 2, 3).number_format.count("0.000000") == 2

    daily = [r for r in rows[summary + 2:] if isinstance(r[0], datetime) and r[1] != "Energy Usage(kWh)"]
    assert len(daily) == 30
    assert daily[0][:2] == [datetime(2026, 9, 1), 6.09]
    assert daily[-1][:2] == [datetime(2026, 9, 30), 16.75645]
    # The daily rows add up to the summary total, read back from the file.
    assert round(sum(r[1] for r in daily), 6) == total

    # The optional header block is present, with the vehicle and source.
    assert rows[0][0] == "Home EV Charging - Novated Lease Claim"
    header = {r[0]: r[1] for r in rows[1:summary] if r[0]}
    assert header["Vehicle"] == "Mitsubishi Outlander PHEV"
    assert header["Period"] == "1 - 30 September 2026"


def test_xlsx_without_header_starts_with_summary(tmp_path: Path) -> None:
    report = build_report("2026-09", SEPTEMBER, 0.407, include_header=False)
    path = tmp_path / "r.xlsx"
    write_xlsx(report, path)
    ws = load_workbook(path).active
    assert ws.cell(1, 2).value == "Energy Usage(kWh)"
    assert ws.cell(2, 2).value == 255.303087


def test_pdf_one_page_for_31_days(tmp_path: Path, monkeypatch) -> None:
    """A 31-day month with every header line still fits on one A4 page."""
    october = {f"2026-10-{d:02d}": 18.007821 for d in range(1, 30)}  # two days missing
    report = build_report("2026-10", october, 0.407, **DETAILS)

    pages: list[int] = []
    original = FPDF.output

    def spy(self, *args, **kwargs):
        pages.append(self.page_no())
        return original(self, *args, **kwargs)

    monkeypatch.setattr(FPDF, "output", spy)
    path = tmp_path / f"{report.file_stem}.pdf"
    write_pdf(report, path)
    assert pages == [1]
    assert path.read_bytes().startswith(b"%PDF")


def test_pdf_golden_text(tmp_path: Path) -> None:
    report = build_report("2026-09", SEPTEMBER, 0.407, **DETAILS)
    path = tmp_path / "r.pdf"
    write_pdf(report, path)
    data = path.read_bytes()
    assert data.startswith(b"%PDF")
    assert b"/Count 1" in data
