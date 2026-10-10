"""Monthly claim report: the numbers, and the XLSX and PDF files.

Kept free of Home Assistant code so it can be tested on its own. The layout
follows the Tapo-app export the lease provider is used to:

    Sep-26     Energy Usage(kWh)   rate        Amount
    2026/09    255.303087          $ 0.407000  $ 103.91

    Sep-26     Energy Usage(kWh)
    2026/09/01  6.09
    ...

with an optional header block above it. Money is calculated with Decimal so
the total and amount are exact (255.303087 x 0.407 = 103.908... -> $103.91,
rounded half-up to cents), never subject to float rounding.
"""

from calendar import monthrange
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

REPORT_TITLE = "Home EV Charging - Novated Lease Claim"
FILE_PREFIX = "Energy_Usage"

CENT = Decimal("0.01")
KWH_PLACES = Decimal("0.000001")

# Number formats copied from the reference spreadsheet.
FMT_MONTH = "mmm-yy"
FMT_PERIOD = "yyyy/mm"
FMT_DATE = "yyyy/mm/dd"
FMT_KWH = "General"
FMT_RATE = '_("$"* #,##0.000000_);_("$"* \\(#,##0.000000\\);_("$"* "-"??_);_(@_)'
FMT_AMOUNT = '_-"$"* #,##0.00_-;\\-"$"* #,##0.00_-;_-"$"* "-"??_-;_-@_-'


@dataclass
class MonthReport:
    """Everything that goes into one month's report."""

    month: str  # "2026-09"
    rate: Decimal
    days: list[tuple[date, Decimal | None]]  # every calendar day; None = no data
    vehicle: str | None = None
    registration: str | None = None
    source: str | None = None
    generated: datetime | None = None
    include_header: bool = True
    notes: list[str] = field(default_factory=list)

    @property
    def first_day(self) -> date:
        return self.days[0][0]

    @property
    def total_kwh(self) -> Decimal:
        """Sum of the daily values, exact to 6 decimals."""
        total = sum((kwh for _, kwh in self.days if kwh is not None), Decimal(0))
        return total.quantize(KWH_PLACES)

    @property
    def amount(self) -> Decimal:
        """Total x rate, rounded half-up to cents."""
        return (self.total_kwh * self.rate).quantize(CENT, rounding=ROUND_HALF_UP)

    @property
    def missing_days(self) -> list[date]:
        return [day for day, kwh in self.days if kwh is None]

    @property
    def label(self) -> str:
        """'Sep-26', the label in the reference's top-left cell."""
        return self.first_day.strftime("%b-%y")

    @property
    def long_name(self) -> str:
        """'September 2026'."""
        return self.first_day.strftime("%B %Y")

    @property
    def file_stem(self) -> str:
        """'Energy_Usage_Sep_2026'."""
        return f"{FILE_PREFIX}_{self.first_day.strftime('%b_%Y')}"


def build_report(
    month: str,
    ledger_days: dict[str, float],
    rate: float,
    **details,
) -> MonthReport:
    """Make a MonthReport from ledger values ({'2026-09-01': 6.09, ...}).

    Every calendar day of the month gets a row; days missing from the ledger
    are kept as None so they can be listed instead of silently becoming 0.
    Values go through str() so 6.09 stays exactly 6.09 in Decimal.
    """
    year, mon = (int(part) for part in month.split("-"))
    days: list[tuple[date, Decimal | None]] = []
    for dom in range(1, monthrange(year, mon)[1] + 1):
        day = date(year, mon, dom)
        kwh = ledger_days.get(day.isoformat())
        days.append((day, None if kwh is None else Decimal(str(kwh))))
    return MonthReport(month=month, rate=Decimal(str(rate)), days=days, **details)


def format_kwh(value: Decimal) -> str:
    """Like Excel's General format: up to 6 decimals, no trailing zeros."""
    text = f"{value.quantize(KWH_PLACES):f}".rstrip("0").rstrip(".")
    return text or "0"


def header_lines(report: MonthReport) -> list[tuple[str, str]]:
    """(label, value) pairs for the optional header block."""
    last_day = report.days[-1][0]
    lines = [("Period", f"{report.first_day.day} - {last_day.day} {report.long_name}")]
    if report.vehicle:
        lines.append(("Vehicle", report.vehicle))
    if report.registration:
        lines.append(("Registration", report.registration))
    if report.source:
        lines.append(("Data source", report.source))
    if report.missing_days:
        lines.append(("Days with no data", ", ".join(f"{d.day} {d:%b}" for d in report.missing_days)))
    if report.generated:
        # Built by hand: "%-d" isn't supported by every C library (e.g. musl).
        g = report.generated
        lines.append(("Generated", f"{g.day} {g:%B %Y %H:%M}"))
    return lines


# --------------------------------------------------------------------- XLSX


def write_xlsx(report: MonthReport, path: Path) -> None:
    """One sheet named like 'Sep-26': header (optional), summary, daily table.

    Numbers are real numbers with number formats (no text, no formulas), so
    the file opens with the values the provider expects.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = report.label
    thin = Side(style="thin")
    boxed = Border(left=thin, right=thin, top=thin, bottom=thin)
    month_start = datetime(report.first_day.year, report.first_day.month, 1)

    row = 1
    if report.include_header:
        ws.cell(row, 1, REPORT_TITLE).font = Font(bold=True, size=13)
        row += 1
        for label, value in header_lines(report):
            ws.cell(row, 1, label).font = Font(bold=True)
            ws.cell(row, 2, value)
            row += 1
        row += 1

    def boxed_row(values: list[tuple[object, str]]) -> None:
        for col, (value, fmt) in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            cell.number_format = fmt
            cell.border = boxed

    # Summary table.
    boxed_row([(month_start, FMT_MONTH), ("Energy Usage(kWh)", "General"), ("rate", "General"), ("Amount", "General")])
    ws.cell(row, 1).alignment = Alignment(horizontal="right")
    row += 1
    boxed_row([
        (month_start, FMT_PERIOD),
        (float(report.total_kwh), FMT_KWH),
        (float(report.rate), FMT_RATE),
        (float(report.amount), FMT_AMOUNT),
    ])
    ws.cell(row, 1).alignment = Alignment(horizontal="left")
    row += 2

    # Daily table.
    boxed_row([(month_start, FMT_MONTH), ("Energy Usage(kWh)", "General")])
    ws.cell(row, 1).alignment = Alignment(horizontal="right")
    row += 1
    for day, kwh in report.days:
        boxed_row([
            (datetime(day.year, day.month, day.day), FMT_DATE),
            (float(kwh) if kwh is not None else "no data", FMT_KWH),
        ])
        ws.cell(row, 1).alignment = Alignment(horizontal="left")
        ws.cell(row, 2).alignment = Alignment(horizontal="right")
        row += 1

    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 19
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 12
    wb.save(path)


# ---------------------------------------------------------------------- PDF


def _pdf_text(text: str) -> str:
    """The built-in PDF fonts only cover Latin-1; swap anything else."""
    text = text.replace("—", "-").replace("–", "-").replace("×", "x")
    return text.encode("latin-1", "replace").decode("latin-1")


def write_pdf(report: MonthReport, path: Path) -> None:
    """A4 portrait, one page: optional header, summary table, daily table."""
    pdf = FPDF(orientation="portrait", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=False)
    pdf.set_margins(20, 18, 20)
    pdf.add_page()
    pdf.set_title(f"{REPORT_TITLE} - {report.long_name}")

    line = 5.2  # row height in mm; 31 days + header still fit on A4

    if report.include_header:
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 8, _pdf_text(REPORT_TITLE), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
        for label, value in header_lines(report):
            pdf.set_font("Helvetica", "B", 9)
            pdf.cell(32, line, _pdf_text(label))
            pdf.set_font("Helvetica", "", 9)
            pdf.multi_cell(0, line, _pdf_text(value), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)

    pdf.set_font("Helvetica", "", 10)

    def table_row(cells: list[tuple[str, float, str]], bold: bool = False) -> None:
        pdf.set_font("Helvetica", "B" if bold else "", 10)
        for text, width, align in cells:
            pdf.cell(width, line, _pdf_text(text), border=1, align=align)
        pdf.ln(line)

    widths = (24, 38, 28, 26)
    table_row([
        (report.label, widths[0], "R"),
        ("Energy Usage(kWh)", widths[1], "L"),
        ("rate", widths[2], "L"),
        ("Amount", widths[3], "L"),
    ], bold=True)
    table_row([
        (f"{report.first_day:%Y/%m}", widths[0], "L"),
        (f"{report.total_kwh:f}", widths[1], "R"),
        (f"$ {report.rate:.6f}", widths[2], "R"),
        (f"$ {report.amount:,.2f}", widths[3], "R"),
    ])
    pdf.ln(line)

    table_row([(report.label, widths[0], "R"), ("Energy Usage(kWh)", widths[1], "L")], bold=True)
    for day, kwh in report.days:
        table_row([
            (f"{day:%Y/%m/%d}", widths[0], "L"),
            (format_kwh(kwh) if kwh is not None else "no data", widths[1], "R"),
        ])

    pdf.output(str(path))
